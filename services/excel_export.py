"""Export data vendor ke Excel dengan header 2 tingkat (kolom NO sejajar dengan header lain)."""
import io
import re
from typing import Dict, List, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.phone import to_whatsapp_url
from core.schema import DATE_COLS, DB_COLS, REVERSE_MAPPING
from core.text_utils import parse_date

_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill("solid", start_color="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_CENTER_PLAIN = Alignment(horizontal="center", vertical="center")

_NUMERIC_COLS = ("iup_countdown", "iup_periode")
_DATE_FORMAT = "DD/MM/YYYY"
_FIRST_DATA_ROW = 3
_LINK_FONT = Font(color="0563C1", underline="single")


def build_excel(df_flat: pd.DataFrame, sheet_name: str = "Data Vendor") -> bytes:
    """DataFrame kolom datar (DB_COLS) -> file .xlsx (bytes)."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = re.sub(r"[\\/*?:\[\]]", "", sheet_name)[:31] or "Data Vendor"

    positions, widths = _write_header(sheet)
    _write_body(sheet, df_flat, positions, widths)
    _apply_layout(sheet, widths)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------ header
def _header_blocks() -> List[Tuple[str, List[Tuple[str, str]]]]:
    """Kelompokkan DB_COLS: kolom tunggal, atau grup berisi beberapa sub-kolom."""
    blocks, i = [], 0
    while i < len(DB_COLS):
        group, sub = REVERSE_MAPPING[DB_COLS[i]]
        if not sub:
            blocks.append((group, [(DB_COLS[i], "")]))
            i += 1
            continue
        members = []
        while i < len(DB_COLS) and REVERSE_MAPPING[DB_COLS[i]][0] == group and REVERSE_MAPPING[DB_COLS[i]][1]:
            members.append((DB_COLS[i], REVERSE_MAPPING[DB_COLS[i]][1]))
            i += 1
        blocks.append((group, members))
    return blocks


def _header_cell(sheet, row: int, col: int, text):
    cell = sheet.cell(row=row, column=col, value=text)
    cell.fill, cell.font, cell.alignment, cell.border = _HEADER_FILL, _HEADER_FONT, _CENTER, _BORDER


def _write_header(sheet) -> Tuple[Dict[str, int], Dict[int, int]]:
    """Tulis 2 baris header. Mengembalikan (posisi kolom per kolom DB, lebar awal per kolom)."""
    _header_cell(sheet, 1, 1, "NO")
    _header_cell(sheet, 2, 1, None)
    sheet.merge_cells(start_row=1, end_row=2, start_column=1, end_column=1)

    positions: Dict[str, int] = {}
    widths: Dict[int, int] = {1: 5}
    col = 2
    for group, members in _header_blocks():
        if len(members) == 1 and not members[0][1]:  # kolom tunggal: gabung 2 baris
            db_col = members[0][0]
            _header_cell(sheet, 1, col, group)
            _header_cell(sheet, 2, col, None)
            sheet.merge_cells(start_row=1, end_row=2, start_column=col, end_column=col)
            positions[db_col] = col
            widths[col] = len(group)
            col += 1
            continue

        for offset, (db_col, sub) in enumerate(members):
            _header_cell(sheet, 1, col + offset, group if offset == 0 else None)
            _header_cell(sheet, 2, col + offset, sub)
            positions[db_col] = col + offset
            widths[col + offset] = len(sub)
        if len(members) > 1:
            sheet.merge_cells(start_row=1, end_row=1, start_column=col, end_column=col + len(members) - 1)
        col += len(members)
    return positions, widths


# ------------------------------------------------------------------ isi data
def _write_body(sheet, df: pd.DataFrame, positions: Dict[str, int], widths: Dict[int, int]) -> None:
    for number, (_, row) in enumerate(df.iterrows(), start=1):
        excel_row = _FIRST_DATA_ROW + number - 1
        no_cell = sheet.cell(row=excel_row, column=1, value=number)
        no_cell.border, no_cell.alignment = _BORDER, _CENTER_PLAIN

        for db_col, col in positions.items():
            raw = row.get(db_col, "")
            text = "" if pd.isna(raw) else str(raw).strip()
            cell = sheet.cell(row=excel_row, column=col)
            _fill_cell(cell, db_col, text)
            cell.border = _BORDER
            widths[col] = max(widths.get(col, 8), min(len(text), 45))


def _fill_cell(cell, db_col: str, text: str) -> None:
    """Isi sel: tanggal sebagai tanggal asli, angka sebagai angka, kontak sebagai link WhatsApp, lainnya teks."""
    if db_col == "no_kontak" and text:
        cell.value = text
        url = to_whatsapp_url(text)
        if url:
            cell.hyperlink = url
            cell.font = _LINK_FONT
        return
    if db_col in DATE_COLS and text:
        parsed = parse_date(text)
        if pd.notna(parsed):
            cell.value = parsed.to_pydatetime()
            cell.number_format = _DATE_FORMAT
            cell.alignment = _CENTER_PLAIN
            return
    if db_col in _NUMERIC_COLS and re.fullmatch(r"-?\d+", text):
        cell.value = int(text)
        cell.alignment = _CENTER_PLAIN
        return
    cell.value = text if text else None


def _apply_layout(sheet, widths: Dict[int, int]) -> None:
    for col, width in widths.items():
        sheet.column_dimensions[get_column_letter(col)].width = max(10, min(width + 3, 48))
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[2].height = 22
    sheet.freeze_panes = "C3"  # NO dan VENDOR tetap terlihat saat digulir
