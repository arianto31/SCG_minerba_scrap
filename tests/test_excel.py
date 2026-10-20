import io
import unittest
from datetime import datetime

import pandas as pd
from openpyxl import Workbook, load_workbook

from core.schema import DB_COLS
from services.excel_export import build_excel
from services.excel_import import parse_vendor_excel

FLAT_HEADER = [
    "NO", "VENDOR", "NAMA PIC", "NO KONTAK", "NIB NOMOR", "NIB TANGGAL", "LOKASI", "SURAT DUKUNGAN NOMOR",
    "SURAT DUKUNGAN TANGGAL", "SURAT DUKUNGAN LOKASI QUARRY", "IUP NOMOR", "IUP NAMA", "IUP TANGGAL",
    "IUP PERIODE", "MINERBAONESMART COMP TANGGAL", "MINERBAONESMART COMP COUNTDOWN", "REMARK",
]


def flat_workbook(rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["Judul laporan"])
    ws.append(FLAT_HEADER)
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class ExcelTest(unittest.TestCase):
    def test_import_flat_header(self):
        data = flat_workbook([
            [1, "ABADI", "Sri", "-", 9120208193905, "30/01/2019", "-", "-", "11/11/2024", "-",
             "570/16", "PT SUMBER", "18/04/2023", 5.116, "16/04/2028", 554, "-"],
            [2, "APIANG", "Apian", "-", 9120311070416, datetime(2019, 10, 1), "-", "-", None, "-",
             "9120", "CV CENTRAL", datetime(2024, 6, 13), 5, datetime(2029, 6, 12), 976, "-"],
        ])
        df = parse_vendor_excel(data).data
        self.assertEqual(len(df), 2)
        self.assertEqual(df.loc[1, "nib_nomor"], "9120208193905")
        self.assertEqual(df.loc[1, "iup_tanggal_terbit"], "2023-04-18")
        self.assertEqual(df.loc[1, "iup_tanggal_berakhir"], "2028-04-16")
        self.assertEqual(df.loc[1, "iup_periode"], "5")          # 5.116 dibulatkan
        self.assertEqual(df.loc[1, "no_kontak"], "")             # '-' dianggap kosong
        self.assertEqual(df.loc[2, "nib_tanggal"], "2019-10-01")
        self.assertEqual(df.loc[1, "iup_countdown"], "")         # dihitung ulang sistem

    def test_golongan_follows_nib_nomor_and_round_trips(self):
        df = pd.DataFrame([{**{c: "" for c in DB_COLS}, "vendor": "ABADI", "nib_nomor": "123",
                            "golongan": "Fly Ash", "nib_tanggal": "2019-01-30"}])
        self.assertEqual(DB_COLS.index("golongan"), DB_COLS.index("nib_nomor") + 1)

        data = build_excel(df, "Vendor")
        sheet = load_workbook(io.BytesIO(data)).active
        row2 = [c.value for c in sheet[2]]
        row1 = [c.value for c in sheet[1]]
        self.assertEqual(row2[row1.index("GOLONGAN")], None)             # kolom tunggal digabung 2 baris
        again = parse_vendor_excel(data).data
        self.assertEqual(again.loc[1, "golongan"], "Fly Ash")
        self.assertEqual(again.loc[1, "nib_tanggal"], "2019-01-30")      # NIB yang terpisah tetap terbaca
        self.assertEqual(again.loc[1, "nib_nomor"], "123")

    def test_flat_header_with_golongan_alias(self):
        wb = Workbook()
        wb.active.append(["VENDOR", "GOLONGAN USAHA", "IUP NAMA"])
        wb.active.append(["ABADI", "Pasir", "PT X"])
        buffer = io.BytesIO()
        wb.save(buffer)
        self.assertEqual(parse_vendor_excel(buffer.getvalue()).data.loc[1, "golongan"], "Pasir")

    def test_contact_becomes_whatsapp_hyperlink_in_excel(self):
        df = pd.DataFrame([{**{c: "" for c in DB_COLS}, "vendor": "A", "no_kontak": "0812-3456-7890"},
                           {**{c: "" for c in DB_COLS}, "vendor": "B", "no_kontak": "Pak Budi"}])
        sheet = load_workbook(io.BytesIO(build_excel(df))).active
        col = [c.value for c in sheet[1]].index("NO KONTAK") + 1
        self.assertEqual(sheet.cell(row=3, column=col).value, "0812-3456-7890")
        self.assertEqual(sheet.cell(row=3, column=col).hyperlink.target, "https://wa.me/6281234567890")
        self.assertIsNone(sheet.cell(row=4, column=col).hyperlink)

    def test_missing_vendor_column_raises(self):
        wb = Workbook()
        wb.active.append(["A", "B"])
        buffer = io.BytesIO()
        wb.save(buffer)
        with self.assertRaises(ValueError):
            parse_vendor_excel(buffer.getvalue())

    def test_export_layout_and_round_trip(self):
        df = pd.DataFrame([{**{c: "" for c in DB_COLS}, "vendor": "ABADI", "iup_nama": "PT X",
                            "iup_tanggal_terbit": "2023-04-18", "iup_tanggal_berakhir": "2028-04-16",
                            "iup_periode": "5", "iup_countdown": "554"}])
        data = build_excel(df, "Vendor Pasir/Batu")
        sheet = load_workbook(io.BytesIO(data)).active

        self.assertEqual(sheet.title, "Vendor PasirBatu")             # karakter terlarang dibuang
        self.assertEqual(sheet["A1"].value, "NO")                    # NO sejajar di baris header pertama
        self.assertIsNone(sheet["A2"].value)                         # tidak turun ke baris kedua
        self.assertEqual(sheet["A3"].value, 1)
        headers = {c.value for c in sheet[2] if c.value}
        self.assertIn("TANGGAL TERBIT", headers)

        again = parse_vendor_excel(data).data
        self.assertEqual(again.loc[1, "iup_tanggal_terbit"], "2023-04-18")
        self.assertEqual(again.loc[1, "iup_tanggal_berakhir"], "2028-04-16")
        self.assertEqual(again.loc[1, "iup_nama"], "PT X")


if __name__ == "__main__":
    unittest.main()
